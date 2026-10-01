"""Extract every tab of the PokeAlliance spreadsheet into data/*.json for the static site.

Each file under data/ owns one subject. Hand-editable settings live in config/*.json
and are never touched here. manifest.json lists every file the page loads; this
script rewrites its "data" list and keeps the "config" list as is.

Usage: python scripts/extract.py
"""
import datetime
import json
import pathlib
import re

import openpyxl

ROOT = pathlib.Path(__file__).resolve().parent.parent
XLSX = ROOT / "Pokedex Pública PokeAlliance (By Mts Vitor).xlsx"
OUT_DIR = ROOT / "data"
MANIFEST = ROOT / "manifest.json"

wb = openpyxl.load_workbook(XLSX, data_only=True)


def clean(v):
    if v is None:
        return None
    if isinstance(v, float):
        return int(v) if v.is_integer() else round(v, 2)
    if isinstance(v, datetime.time):
        return f"{v.hour}:{v.minute:02d}"
    if isinstance(v, str):
        v = v.strip()
        return v or None
    return v


def rows(sheet, start=2, cols=None):
    ws = wb[sheet]
    for r in ws.iter_rows(min_row=start, values_only=True):
        r = [clean(v) for v in r]
        if cols:
            r = r[:cols]
        if any(v not in (None, "-") for v in r):
            yield r


def compact(xs):
    return [x for x in xs if x not in (None, "", "-")]


def url(v):
    if isinstance(v, str) and re.match(r"^(https?://)?(www\.)?(imgur|discord|youtu|tinyurl|wiki\.|docs\.google)", v):
        return v if v.startswith("http") else "https://" + v
    return None


data = {}

# ---- Tier list (main table + curated shiny lists) ----
ws = wb["Tier List"]
tiers = []
for r in rows("Tier List", cols=3):
    if r[0] and r[1]:
        tiers.append({"name": r[0], "tier": r[1], "moveset": r[2]})
data["tiers"] = tiers
highlights = {}
for col, label in (("I", "Ultra Rare"), ("L", "T1"), ("O", "Super Rare"), ("R", "Legendary")):
    vals = [clean(c.value) for c in ws[col]][1:]
    highlights[label] = compact(vals)
data["tierHighlights"] = highlights

# ---- Drops (pokemon -> items) ----
drops = {}
for r in rows("Drops"):
    if r[0]:
        drops[r[0]] = compact(r[1:])
data["drops"] = drops

# ---- Medals ----
data["medals"] = [
    {"name": r[0], "buff": r[1], "debuff": r[2]} for r in rows("Medals", cols=3) if r[0]
]

# ---- Hunt locations ----
data["locations"] = [
    {"name": r[0], "wild": url(r[1]), "normal": url(r[2]), "hoenn": url(r[3])}
    for r in rows("Localizações", cols=4)
    if r[0]
]

# ---- Tasks (NPC tasks per pokemon) ----
tasks = []
for r in rows("Tasks", cols=7):
    npcs = [{"npc": r[i], "url": url(r[i + 1])} for i in (1, 3, 5) if r[i]]
    if r[0]:
        tasks.append({"name": r[0], "npcs": npcs})
data["tasks"] = tasks

# ---- Linked tasks ----
linked, remaining = [], []
section = "main"
for r in rows("Linked Tasks", cols=5):
    if r[1] in ("Restantes", "Restantes:"):
        section = "summary" if r[1] == "Restantes" else "remaining"
        continue
    if section == "main" and isinstance(r[0], (int, float)):
        linked.append({"qty": r[0], "name": r[1], "type": r[2], "url": url(r[3]), "kph": r[4]})
    elif section == "remaining" and r[1]:
        remaining.append({"name": r[1], "tier": r[2], "where": r[3]})
data["linked"] = linked
data["linkedRemaining"] = remaining

# ---- Talents ----
data["talents"] = [
    {"item": r[0], "source": r[1], "qty": r[3], "type": r[4], "level": r[5], "buff": r[7]}
    for r in rows("PokeTalents", cols=9)
    if r[0] and r[4]
]

# ---- Boost items per type + fragment table ----
data["boost"] = {r[0]: compact(r[1:]) for r in rows("Boost", start=1) if r[0]}
ws = wb["Search Boost Items"]
fix = {datetime.datetime(2025, 6, 10): "6-10", datetime.datetime(2025, 11, 15): "11-15"}
levels = [fix.get(c.value, c.value) for c in ws["E3":"N3"][0]]


def frag_row(first, last):
    return [clean(c.value) for c in ws[f"{first}4":f"{last}4"][0]]


data["fragments"] = {
    "levels": levels,
    "min": frag_row("E", "N"),
    "avg": frag_row("P", "Y"),
    "max": frag_row("AA", "AJ"),
    "days": 477,
}

# ---- Dungeons (dens) ----
dg_items = {r[0]: {"players": r[1], "drops": compact(r[2:])} for r in rows("DgItems") if r[0]}
dens = []
for r in rows("DgMobs", cols=11):
    if not r[0]:
        continue
    item = dg_items.get(r[0], {})
    dens.append({
        "name": r[0], "players": r[1], "mobsCount": r[2], "xp": r[3], "time": r[4],
        "mobs": compact(r[5:10]), "xph": r[10], "drops": item.get("drops", []),
    })
for name, item in dg_items.items():
    if not any(d["name"] == name for d in dens):
        dens.append({"name": name, "players": item["players"], "drops": item["drops"], "mobs": []})
data["dens"] = dens

dungeons, disabled, section = [], [], "on"
for r in rows("Dungeons", cols=5):
    if r[0] == "Desativados":
        section = "off"
        continue
    if section == "on":
        dungeons.append({"name": r[0], "url": url(r[1]), "hunts": compact(r[2:])})
    elif r[0]:
        disabled.append(r[0])
data["dungeons"] = dungeons
data["dungeonsDisabled"] = disabled

# ---- Gym ----
ws = wb["GYM"]
data["gym"] = [
    {"city": r[0], "task1": r[1], "task2": r[2], "url": url(r[3])}
    for r in ws.iter_rows(min_row=2, max_row=9, max_col=4, values_only=True)
]
cities = [c.value for c in ws[12][:8]]
data["gymLeaders"] = {
    city: [clean(ws.cell(row=i, column=j + 1).value) for i in range(13, 19)]
    for j, city in enumerate(cities)
}
data["gymInfo"] = clean(ws["E1"].value)


# ---- Rocket / Police (4 blocks per band, npc + recommendation) ----
def teams(sheet, bands):
    ws = wb[sheet]
    out = []
    for head_row, first, last in bands:
        for col in (1, 4, 7, 10):
            name = clean(ws.cell(row=head_row, column=col).value)
            if not name:
                continue
            members = []
            for rr in range(first, last + 1):
                a = clean(ws.cell(row=rr, column=col).value)
                b = clean(ws.cell(row=rr, column=col + 1).value)
                if a:
                    members.append({"npc": a, "counter": b})
            out.append({"name": name, "members": members})
    return out


data["rocket"] = teams("Rocket", [(1, 2, 7), (9, 10, 15), (17, 18, 23)])
data["rocket"].append({
    "name": "GIOVANNI",
    "note": clean(wb["Rocket"]["C25"].value),
    "members": [
        {"npc": clean(wb["Rocket"].cell(row=r, column=1).value), "counter": clean(wb["Rocket"].cell(row=r, column=2).value)}
        for r in range(26, 32)
    ],
})
data["rocketInfo"] = clean(wb["Rocket"]["L1"].value)
data["police"] = teams("Police", [(1, 2, 7), (9, 10, 15), (17, 18, 23)])
data["policeInfo"] = clean(wb["Police"]["L1"].value)

# ---- Hazard ----
data["hazard"] = [
    {"npc": r[0], "url": url(r[1]), "task": r[2]} for r in rows("Hazard Tasks", cols=3) if r[0]
]

# ---- Reference tables ----
data["brokes"] = [{"tier": r[0], "max": r[1]} for r in rows("Brokes", cols=2) if r[0]]
data["brokesInfo"] = clean(wb["Brokes"]["D1"].value)

ws = wb["Shiny Rate"]
rate_cols = []
for c in range(1, 15, 2):
    rate_cols.append({
        "label": clean(ws.cell(row=1, column=c).value),
        "values": {clean(ws.cell(row=r, column=c).value): clean(ws.cell(row=r, column=c + 1).value) for r in range(3, 8)},
    })
data["shinyRate"] = rate_cols
data["shinyRateForm"] = clean(ws["A10"].value)

ws = wb["Star Level"]
data["starLevel"] = {
    clean(ws.cell(row=r, column=1).value): [clean(ws.cell(row=r, column=c).value) for c in range(2, 12)]
    for r in range(3, 9)
}
data["starInfo"] = clean(wb["Star"]["A4"].value)

ws = wb["Runes"]
runes = []
for c in range(2, 18, 2):
    name = clean(ws.cell(row=1, column=c).value)
    lv = []
    for r in range(3, 8):
        p, b = clean(ws.cell(row=r, column=c).value), clean(ws.cell(row=r, column=c + 1).value)
        if p is not None or b is not None:
            lv.append({"level": clean(ws.cell(row=r, column=1).value), "points": p, "bonus": b})
    runes.append({"name": name, "levels": lv})
data["runes"] = runes

ws = wb["Damage"]
data["damage"] = {
    "tiers": [clean(c.value) for c in ws[1][1:8]],
    "rows": [[clean(c.value) for c in ws[r][:8]] for r in range(2, 6)],
    "note": clean(ws["A7"].value),
}

# ---- Guides / text ----
data["faq"] = [{"keys": str(r[0]), "text": r[1]} for r in rows("FAQ", start=1, cols=2) if r[0] and r[1]]
data["porygon"] = [{"title": r[0].rstrip(":"), "text": r[1]} for r in rows("Porygon", start=1, cols=2) if r[0]]
data["bh"] = clean(wb["BH"]["A1"].value)

# One file per subject; the keys inside each file are what the page reads.
FILES = {
    "tiers.json": ["tiers", "tierHighlights"],
    "drops.json": ["drops"],
    "medals.json": ["medals"],
    "hunt-locations.json": ["locations"],
    "npc-tasks.json": ["tasks"],
    "linked-tasks.json": ["linked", "linkedRemaining"],
    "hazard-tasks.json": ["hazard"],
    "gym.json": ["gym", "gymLeaders", "gymInfo"],
    "bounty-hunter.json": ["bh"],
    "talents.json": ["talents"],
    "boost.json": ["boost", "fragments"],
    "dens.json": ["dens"],
    "dungeons.json": ["dungeons", "dungeonsDisabled"],
    "porygon.json": ["porygon"],
    "rocket.json": ["rocket", "rocketInfo"],
    "police.json": ["police", "policeInfo"],
    "star.json": ["starLevel", "starInfo"],
    "brokes.json": ["brokes", "brokesInfo"],
    "shiny-rate.json": ["shinyRate", "shinyRateForm"],
    "runes.json": ["runes"],
    "damage.json": ["damage"],
    "faq.json": ["faq"],
}
# Hand-written files in data/ that don't come from the spreadsheet: kept as is and listed in the manifest
MANUAL = ["hoenn-teams.json"]

assigned = [k for keys in FILES.values() for k in keys]
assert sorted(assigned) == sorted(data), f"unassigned keys: {set(data) ^ set(assigned)}"

OUT_DIR.mkdir(exist_ok=True)
for old_file in OUT_DIR.glob("*.json"):
    if old_file.name not in FILES and old_file.name not in MANUAL:
        old_file.unlink()
for name, keys in FILES.items():
    path = OUT_DIR / name
    path.write_text(json.dumps({k: data[k] for k in keys}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"  data/{name}: {', '.join(keys)} ({path.stat().st_size:,} bytes)")

manifest = json.loads(MANIFEST.read_text(encoding="utf-8")) if MANIFEST.exists() else {"config": []}
manifest["data"] = [f"data/{name}" for name in [*FILES, *MANUAL]]
MANIFEST.write_text(json.dumps({"data": manifest["data"], "config": manifest.get("config", [])}, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
print("wrote", MANIFEST.name)

# Keep bundle.js (used when index.html is opened from disk) in sync with the new data
import build  # noqa: E402  (scripts/build.py, same folder)

build.build()
